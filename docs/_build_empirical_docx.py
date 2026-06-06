# -*- coding: utf-8 -*-
"""Build the empirical-part academic DOCX (French)."""
import os
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(ROOT, "docs", "_empirical_figures")
RFIG = os.path.join(ROOT, "reports", "figures")
OUT = os.path.join(ROOT, "docs", "Partie_Empirique_Methodologie_et_Resultats.docx")

# Palette
INK = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
SLATE = RGBColor(0x33, 0x41, 0x55)
GREY = RGBColor(0x47, 0x55, 0x69)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()

# ---- base styles ----
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.font.color.rgb = INK
pf = normal.paragraph_format
pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
pf.line_spacing = 1.25
pf.space_after = Pt(6)
pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _set_heading(style_id, size, color, before, after):
    st = doc.styles[style_id]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True


_set_heading("Heading 1", 17, BLUE, 18, 8)
_set_heading("Heading 2", 14, INK, 14, 6)
_set_heading("Heading 3", 12, SLATE, 10, 4)

# ---- page setup: Letter, 1" margins ----
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, side, Inches(1))
CONTENT_W = 6.5


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def cell_text(cell, text, bold=False, color=None, size=9.5, align=None, white=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    if align:
        p.alignment = align
    r = p.add_run(str(text))
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    if white:
        r.font.color.rgb = WHITE
    elif color:
        r.font.color.rgb = color
    else:
        r.font.color.rgb = INK


def add_table(headers, rows, widths=None, header_fill="1D4ED8",
              first_col_bold=True, zebra="F1F5F9", font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        cell_text(hdr[i], h, bold=True, white=True, size=font,
                  align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(hdr[i], header_fill)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cell_text(cells[i], val, bold=(first_col_bold and i == 0), size=font,
                      align=WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            if zebra and ridx % 2 == 1:
                shade(cells[i], zebra)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


_fig_counter = {"n": 0}


def add_figure(path, caption, max_w=6.2):
    if not os.path.exists(path):
        print("MISSING FIGURE:", path)
        return
    with Image.open(path) as im:
        w, h = im.size
    disp_w = min(max_w, CONTENT_W)
    disp_h = disp_w * h / w
    if disp_h > 7.0:
        disp_h = 7.0
        disp_w = disp_h * w / h
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(path, width=Inches(disp_w))
    _fig_counter["n"] += 1
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(f"Figure {_fig_counter['n']} — {caption}")
    r.font.name = "Arial"
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = GREY


def para(text, italic=False, space_after=6, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.font.italic = italic
    r.font.bold = bold
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r = p.add_run(bold_lead)
        r.font.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def h1(t): return doc.add_heading(t, level=1)
def h2(t): return doc.add_heading(t, level=2)
def h3(t): return doc.add_heading(t, level=3)


def page_break():
    doc.add_page_break()


# =====================================================================
# TITLE PAGE
# =====================================================================
for _ in range(3):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("PARTIE EMPIRIQUE")
r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = BLUE; r.font.name = "Arial"

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Pipeline, méthodologie, simulation et résultats")
r.font.size = Pt(16); r.font.color.rgb = SLATE; r.font.name = "Arial"

doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Système de détection de la fraude par carte bancaire\n"
              "en temps réel avec intelligence artificielle explicable")
r.font.size = Pt(13); r.font.italic = True; r.font.color.rgb = INK; r.font.name = "Arial"

for _ in range(2):
    doc.add_paragraph()

# thin rule
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
pPr = p._p.get_or_add_pPr()
pbdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "12")
bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "2563EB")
pbdr.append(bottom); pPr.append(pbdr)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Projet de Fin d'Études (PFE)")
r.font.size = Pt(12); r.font.bold = True; r.font.name = "Arial"
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Document méthodologique et expérimental")
r.font.size = Pt(11); r.font.color.rgb = GREY; r.font.name = "Arial"

page_break()

# =====================================================================
# TABLE OF CONTENTS
# =====================================================================
h1("Table des matières")
tocp = doc.add_paragraph()
run = tocp.add_run()
fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
instr.text = 'TOC \\o "1-3" \\h \\z \\u'
fld_sep = OxmlElement("w:fldChar"); fld_sep.set(qn("w:fldCharType"), "separate")
fld_t = OxmlElement("w:t")
fld_t.text = "Mettez à jour ce champ dans Word (clic droit ▸ Mettre à jour les champs, ou F9)."
fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
run._r.append(fld_begin); run._r.append(instr); run._r.append(fld_sep)
run._r.append(fld_t); run._r.append(fld_end)

page_break()

# =====================================================================
# 1. INTRODUCTION
# =====================================================================
h1("1. Introduction à la partie empirique")
para(
    "Cette partie empirique décrit, de manière reproductible et détaillée, la "
    "démarche expérimentale mise en œuvre pour concevoir, entraîner, évaluer et "
    "déployer un système de détection de la fraude par carte bancaire fonctionnant "
    "en temps réel. L'objectif n'est pas seulement d'obtenir un classifieur "
    "performant, mais de construire une chaîne de traitement complète — depuis la "
    "donnée brute jusqu'à la décision explicable présentée à l'utilisateur — "
    "respectant les contraintes méthodologiques (absence de fuite temporelle), "
    "économiques (coût asymétrique des erreurs) et réglementaires (explicabilité "
    "des décisions automatisées)."
)
para(
    "Le système repose sur deux volets complémentaires et interconnectés. Le "
    "premier volet est un pipeline d'apprentissage supervisé hors-ligne, entraîné "
    "et validé sur un jeu de données réel de transactions par carte. Le second "
    "volet est une simulation bancaire en temps réel : un générateur de "
    "transactions synthétiques alimente un moteur de scoring dynamique et "
    "explicable, dont les décisions sont diffusées en continu vers un tableau de "
    "bord et persistées en base pour l'analyse décisionnelle. Les deux volets "
    "partagent une même infrastructure (API, base de données, couche de "
    "restitution), ce qui assure la cohérence de bout en bout."
)
para(
    "Les sections suivantes présentent successivement : la vue d'ensemble de "
    "l'architecture et des relations entre composants (section 2) ; le pipeline "
    "d'apprentissage hors-ligne et sa méthodologie (section 3) ; la simulation "
    "temps réel, sa base de génération et son moteur de scoring (section 4) ; "
    "l'architecture technique réellement développée — conteneurs Docker, backend "
    "FastAPI, base PostgreSQL, frontend et intégration Power BI (section 5) ; le "
    "protocole d'évaluation et les indicateurs retenus (section 6) ; la "
    "restitution décisionnelle dans le tableau de bord Power BI, page par page "
    "(section 7) ; puis une discussion critique assortie des perspectives "
    "(section 8) et une synthèse (section 9)."
)
para(
    "Conformément à une exigence de rigueur, cette partie ne décrit que ce qui a "
    "été réellement implémenté : les méthodes effectivement appliquées, "
    "l'architecture réellement développée, les variables effectivement créées et "
    "le protocole réellement utilisé. Les fonctionnalités envisagées mais non "
    "encore réalisées sont explicitement signalées comme perspectives ou éléments "
    "en cours de développement. Enfin, les résultats chiffrés définitifs "
    "(performances, coûts, SHAP, simulation) seront intégrés à la section 6 sur la "
    "base des tableaux, métriques et exports finaux ; la présente version en "
    "définit le protocole de production et d'analyse."
)

# =====================================================================
# 2. VUE D'ENSEMBLE
# =====================================================================
h1("2. Vue d'ensemble : architecture et relations entre composants")
para(
    "L'architecture générale est organisée autour de deux volets reliés par un "
    "noyau partagé. La Figure 1 en donne la représentation synthétique : à gauche, "
    "le volet d'apprentissage supervisé produit des modèles sérialisés ; à droite, "
    "le volet de simulation temps réel génère et score des transactions ; au "
    "centre, le magasin de modèles (models_store/), l'API FastAPI et la couche de "
    "persistance ORM assurent la liaison ; en bas, le tableau de bord et Power BI "
    "constituent la couche de restitution."
)
add_figure(os.path.join(FIG, "fig1_architecture.png"),
           "Architecture générale du système et relations entre les deux volets.")

h2("2.1. Les éléments du système et leurs rôles")
add_table(
    ["Composant", "Rôle dans la chaîne empirique"],
    [
        ["Jeu de données creditcard.csv", "Source réelle pour l'apprentissage supervisé hors-ligne."],
        ["Pipeline d'entraînement (scripts/)", "EDA, prétraitement, feature engineering, entraînement, évaluation."],
        ["models_store/", "Artefacts persistés : modèles, préprocesseur, ingénieur de variables."],
        ["Générateur synthétique (src/simulation)", "Crée clients et transactions réalistes pour la démonstration temps réel."],
        ["Moteur de scoring (src/scoring)", "Évalue chaque transaction et produit une décision explicable."],
        ["Service de streaming (src/streaming)", "Boucle asynchrone, broker d'événements, diffusion SSE/WebSocket."],
        ["API FastAPI (src/api)", "Expose la prédiction, l'explication et les compteurs agrégés."],
        ["Base de données (src/database)", "ORM SQLAlchemy ; SQLite en développement, PostgreSQL en production."],
        ["Tableau de bord + Power BI", "Restitution temps réel et analyse décisionnelle a posteriori."],
    ],
    widths=[2.3, 4.2],
)

h2("2.2. Relations et flux de données")
para(
    "Les relations entre composants suivent deux chemins. Le chemin "
    "« apprentissage » est descendant et ponctuel : la donnée brute traverse le "
    "pipeline, qui dépose ses artefacts dans models_store/. Le chemin "
    "« inférence » est continu : le générateur produit une transaction, le moteur "
    "de scoring la note, la décision est persistée par l'ORM puis publiée par le "
    "broker vers les abonnés (tableau de bord). L'API joue le rôle de point "
    "d'entrée unique, tandis que la base de données constitue la mémoire commune "
    "que Power BI interroge pour l'analyse. Cette séparation claire entre la phase "
    "d'apprentissage et la phase d'inférence est une garantie méthodologique : le "
    "modèle entraîné et figé ne dépend jamais des données qu'il rencontrera en "
    "production."
)

page_break()

# =====================================================================
# 3. VOLET 1 — PIPELINE HORS-LIGNE
# =====================================================================
h1("3. Volet 1 — Pipeline d'apprentissage supervisé")
para(
    "Le premier volet est une chaîne de traitement séquentielle, illustrée par la "
    "Figure 2, dont chaque étape produit des sorties consommées par la suivante. "
    "La conception privilégie trois principes : la reproductibilité (graine "
    "aléatoire fixe, artefacts persistés), l'absence de fuite temporelle, et "
    "l'orientation économique de la décision finale."
)
add_figure(os.path.join(FIG, "fig2_pipeline_offline.png"),
           "Pipeline d'apprentissage supervisé : enchaînement des dix étapes et leurs sorties.")

h2("3.1. Le jeu de données")
para(
    "Le pipeline est entraîné sur le jeu « European Credit Card Fraud Detection » "
    "(Pozzolo et al., 2015), qui rassemble 284 807 transactions de porteurs "
    "européens collectées sur deux jours de septembre 2013, dont seulement 492 "
    "frauduleuses, soit un taux de 0,17 %. Vingt-huit des trente variables "
    "d'entrée (V1 à V28) sont issues d'une analyse en composantes principales "
    "(ACP) anonymisante ; seules les variables Time et Amount restent "
    "interprétables. Ce déséquilibre extrême et cette anonymisation conditionnent "
    "l'ensemble des choix méthodologiques qui suivent."
)
add_table(
    ["Caractéristique", "Valeur"],
    [
        ["Transactions totales", "284 807"],
        ["Transactions frauduleuses", "492 (0,17 %)"],
        ["Variables d'entrée", "30 (V1–V28 issues d'une ACP + Time + Amount)"],
        ["Variable cible", "Class (0 = légitime, 1 = fraude)"],
        ["Fenêtre d'observation", "48 heures"],
        ["Jeu de test (hold-out chronologique)", "≈ 56 962 transactions, dont 75 fraudes"],
    ],
    widths=[3.0, 3.5],
)

h2("3.2. Validation temporelle anti-fuite")
para(
    "Un découpage aléatoire entraînement/test permettrait au modèle de « regarder "
    "dans le futur » ; sur 48 heures de données, même une faible fuite gonfle les "
    "métriques de plusieurs points de PR-AUC. Trois protections sont donc "
    "appliquées de façon stricte :"
)
bullet("un découpage chronologique pour l'évaluation finale (les 20 % les plus "
       "récents servent de hold-out, jamais vus à l'entraînement) ;", "Hold-out chronologique : ")
bullet("une validation croisée par TimeSeriesSplit (5 plis) : chaque pli prédit "
       "le bloc chronologique suivant, jamais un bloc passé ;", "TimeSeriesSplit : ")
bullet("l'ajustement de tous les préprocesseurs (mise à l'échelle, statistiques "
       "glissantes, z-scores) sur le seul pli d'entraînement.", "Ajustement sur train uniquement : ")

h2("3.3. Ingénierie des variables")
para(
    "Sur la base des deux variables interprétables et des composantes ACP, une "
    "couche d'ingénierie enrichit la représentation : logarithme du montant, "
    "indicateur de transaction nocturne, heure de la journée, statistiques "
    "glissantes (fenêtres de 10, 50 et 200 transactions), z-score du montant, "
    "indicateurs de vélocité et termes d'interaction. Chaque transformation est "
    "calculée de manière causale, c'est-à-dire sans utiliser d'information "
    "postérieure à la transaction considérée."
)

h2("3.4. Traitement du déséquilibre de classes")
para(
    "Avec 0,17 % de positifs, un classifieur naïf atteindrait 99,83 % d'exactitude "
    "en prédisant toujours « légitime ». Plusieurs stratégies de rééquilibrage "
    "sont donc comparées : aucune, SMOTE, BorderlineSMOTE, ADASYN, SMOTEENN et "
    "pondération des classes (class_weight). La stratégie retenue par défaut est "
    "SMOTEENN — une approche hybride combinant sur-échantillonnage et nettoyage "
    "des exemples ambigus — tandis que la pondération des classes reste "
    "recommandée en production car elle n'introduit aucune donnée synthétique, "
    "ce qui facilite l'audit réglementaire."
)
add_table(
    ["Stratégie", "Rappel", "Précision", "Remarque"],
    [
        ["Aucune", "Faible", "Élevée", "Nombreuses fraudes manquées (coût FN élevé)."],
        ["Sous-échantillonnage", "Moyen", "Faible", "Perte de signal, hausse des faux positifs."],
        ["SMOTE", "Élevé", "Moyenne", "Bruit synthétique aux frontières."],
        ["BorderlineSMOTE", "Élevé", "Moy.-élevée", "Meilleure variante SMOTE classique."],
        ["ADASYN", "Élevé", "Moyenne", "Agressif sur les exemples difficiles."],
        ["SMOTEENN", "Élevé", "Élevée", "Défaut recommandé — nettoyage hybride."],
        ["class_weight", "Moy.-élevé", "Élevée", "Sans manipulation de données — sûr en production."],
    ],
    widths=[1.7, 1.0, 1.1, 2.7], font=9,
)

h2("3.5. La famille de modèles")
para(
    "Sept modèles sont entraînés et comparés, de la régression logistique de "
    "référence jusqu'aux ensembles avancés : régression logistique, forêt "
    "aléatoire, XGBoost, LightGBM, CatBoost, un ensemble par vote pondéré "
    "(voting, poids 1/2/3/2 pour LR/RF/XGB/LGBM) et un ensemble par empilement "
    "(stacking, méta-apprenant logistique). Cette progression permet de mesurer "
    "le gain apporté par les méthodes d'ensemble par rapport aux modèles isolés."
)

h2("3.6. Optimisation des hyper-paramètres (Optuna)")
para(
    "Les hyper-paramètres sont optimisés par recherche bayésienne avec Optuna "
    "(échantillonneur TPE, élagueur médian, 30 essais, métrique cible PR-AUC). "
    "L'optimisation est encapsulée dans la validation temporelle afin que la "
    "sélection des hyper-paramètres elle-même ne provoque aucune fuite."
)

h2("3.7. Calibration et optimisation du seuil de décision")
para(
    "La probabilité produite par chaque modèle est calibrée par régression "
    "isotonique, de sorte qu'un score de 0,8 corresponde effectivement à une "
    "fréquence de fraude de 80 %. Le seuil de décision n'est pas fixé "
    "arbitrairement à 0,5 : il est optimisé selon trois critères (F1, MCC, et "
    "coût économique). Une décision en deux temps est également introduite — "
    "APPROUVER / EXAMINER / BLOQUER — la bande « examiner » orientant les cas "
    "ambigus vers un analyste humain plutôt que vers un blocage automatique."
)

h2("3.8. La fonction de coût asymétrique")
para(
    "Le cœur de l'orientation économique est une fonction de coût asymétrique "
    "C(s) = FN(s)·C_FN + FP(s)·C_FP évaluée sur toute la grille de seuils s. Le "
    "mode retenu est « sensible au montant » : le coût d'une fraude non détectée "
    "dépend du montant réellement en jeu, augmenté de pénalités fixes, tandis que "
    "le coût d'un blocage légitime agrège les frais d'investigation, la marge "
    "perdue et l'insatisfaction client."
)
add_table(
    ["Composante de coût", "Valeur", "Interprétation"],
    [
        ["C_FN = montant × loss_factor", "× 1,0", "Montant frauduleux remboursé."],
        ["  + pénalité de chargeback", "25,0", "Pénalité réseau/scheme fixe."],
        ["  + coût de traitement", "18,0", "Coût opérationnel de gestion."],
        ["  + proxy réputationnel", "10,0", "Coût réputationnel agrégé."],
        ["C_FP = investigation", "8,0", "Coût d'analyse d'un faux positif."],
        ["  + marge perdue", "3,0", "Revenu perdu sur la transaction bloquée."],
        ["  + insatisfaction", "1,5", "Proxy d'insatisfaction client (= 12,5 au total)."],
    ],
    widths=[2.7, 1.0, 2.8], font=9,
)
para(
    "Cette fonction permet de comparer chaque seuil optimal à un seuil de "
    "référence (0,50) et de chiffrer, en euros, l'économie réalisée — un résultat "
    "directement exploitable par une direction des risques (voir section 6.4).",
    italic=True,
)

page_break()

# =====================================================================
# 4. VOLET 2 — SIMULATION TEMPS RÉEL
# =====================================================================
h1("4. Volet 2 — La simulation bancaire en temps réel")
para(
    "Le second volet répond à une question concrète : comment démontrer, en "
    "séance et de façon réaliste, le fonctionnement d'un moteur antifraude "
    "bancaire ? Le jeu de données européen étant anonymisé (composantes ACP "
    "dépourvues de sémantique) et limité à 48 heures, il ne permet ni de "
    "rejouer des scénarios de fraude intelligibles, ni d'exposer des signaux "
    "métier (marchand, géolocalisation, appareil). La simulation comble ce "
    "manque en générant des transactions synthétiques riches et explicables."
)

h2("4.1. Sur quoi la simulation est fondée")
para(
    "La simulation est entièrement déterministe et reproductible : elle repose sur "
    "un générateur pseudo-aléatoire de graine fixe (42). Elle s'appuie sur des "
    "catalogues métier réalistes et sur une modélisation à deux niveaux — le "
    "client, puis la transaction — qui reproduit la logique d'un système bancaire "
    "réel. Les briques de base sont les suivantes :"
)
bullet("1 000 clients synthétiques dotés d'un profil de risque, d'un pays "
       "d'origine, d'un comportement type, d'une dépense moyenne, d'une ancienneté "
       "de compte et d'un historique de fraude ;", "Base clientèle : ")
bullet("37 pays avec leurs villes et coordonnées géographiques réelles, leurs "
       "devises, ainsi que des catalogues de marchands, de catégories, de "
       "navigateurs, de systèmes d'exploitation et de types de carte ;", "Catalogues géographiques et marchands : ")
bullet("dix scénarios de fraude injectables (montant anormal, transaction "
       "nocturne, pays étranger, voyage impossible, rafale de transactions, carte "
       "volée, fraude robotisée, multiplication d'adresses IP, prise de contrôle "
       "de compte, dépense anormale).", "Scénarios de fraude : ")

h3("4.1.1. Génération des clients")
para(
    "Chaque client est tiré selon une distribution réaliste de profils de risque. "
    "Les attributs (dépense moyenne, ancienneté, historique de fraude, comportement "
    "type, appareils connus, pays habituels) sont conditionnés par ce profil : un "
    "client VIP dépense davantage et voyage, un fraudeur est récent et présente un "
    "historique de fraude, etc."
)
add_table(
    ["Profil", "Part", "Dépense moy.", "Comportement type"],
    [
        ["LOW", "54 %", "15–160", "Profils de détail quotidien."],
        ["MEDIUM", "24 %", "45–280", "Détail / ménage."],
        ["VIP", "12 %", "350–1 800", "Voyageur, chef d'entreprise."],
        ["HIGH", "8 %", "90–650", "Digital, voyageur."],
        ["FRAUDSTER", "2 %", "80–350", "Cartes-cadeaux, crypto, jeu."],
    ],
    widths=[1.6, 0.9, 1.4, 2.6], font=9,
)

h3("4.1.2. Génération des transactions et scénarios")
para(
    "Pour chaque transaction, un client est tiré au hasard, puis un scénario est "
    "choisi : un fraudeur déclenche un scénario frauduleux dans 55 % des cas, "
    "tandis qu'un client ordinaire suit le taux de fraude paramétré (8 % par "
    "défaut). Le scénario gouverne ensuite tous les attributs de la transaction — "
    "le montant (multiplicateur de la dépense moyenne), la localisation (pays "
    "étranger pour un voyage impossible), l'intelligence IP (botnet, proxy, Tor), "
    "l'appareil (connu ou non), le marchand (cartes-cadeaux et crypto pour les "
    "scénarios de cash-out) — de manière à produire un signal frauduleux cohérent "
    "et intelligible."
)

h2("4.2. Le moteur de scoring dynamique explicable")
para(
    "Le moteur de scoring est le cœur du volet temps réel. Contrairement à un "
    "modèle « boîte noire », il est volontairement additif et inspectable : chaque "
    "signal de risque ajoute un nombre de points motivé par une raison en langage "
    "naturel. La somme des contributions, bornée à l'intervalle 0–100, constitue "
    "le score de risque ; une transformation sigmoïde en déduit une probabilité de "
    "fraude. La Figure 3 schématise ce mécanisme."
)
add_figure(os.path.join(FIG, "fig3_scoring_engine.png"),
           "Moteur de scoring dynamique : des signaux aux contributions additives, "
           "puis au score, à la probabilité et à la décision.")

para(
    "Les signaux pris en compte couvrent le profil client, le montant absolu et "
    "relatif (ratio au montant moyen du client), l'heure, le risque géographique "
    "et le caractère étranger du pays, la réputation de l'adresse IP (VPN, proxy, "
    "Tor, botnet, liste noire), l'appareil inconnu, la catégorie de marchand et le "
    "type de carte, l'ancienneté et l'historique de fraude. À ces signaux "
    "statiques s'ajoutent trois signaux comportementaux calculés sur l'historique "
    "récent du client :"
)
bullet("nombre de transactions dans les 60 dernières secondes (détection de "
       "rafale) ;", "Vélocité : ")
bullet("distance et temps écoulé entre deux transactions, via la formule de "
       "haversine ; au-delà de 500 km à plus de 900 km/h, le déplacement est jugé "
       "physiquement impossible ;", "Voyage impossible : ")
bullet("nombre d'adresses IP distinctes observées pour un même client sur une "
       "fenêtre de 300 secondes.", "Rotation d'IP : ")

para(
    "Quelques pondérations de risque illustrent la logique additive du moteur "
    "(échelle 0–100) :"
)
add_table(
    ["Signal", "Points de risque (extraits)"],
    [
        ["Profil client (base)", "LOW 6 · MEDIUM 16 · HIGH 32 · FRAUDSTER 70"],
        ["Montant absolu", "≥ 1 500 : 9 · ≥ 5 000 : 20 · ≥ 10 000 : 28"],
        ["Montant vs moyenne client", "× 2,5 : 10 · × 5 : 20 · × 10 : 30"],
        ["Réputation IP", "vpn 10 · proxy 14 · tor 22 · botnet 28 · liste noire 35"],
        ["Voyage impossible", "22 (ou 32 si vitesse > 1 200 km/h)"],
        ["Rafale (60 s)", "≥ 4 : 9 · ≥ 8 : 18 · ≥ 15 : 28"],
        ["Scénario injecté", "fraude robotisée 32 · prise de contrôle 38 · carte volée 36"],
    ],
    widths=[2.4, 4.1], font=9,
)

h2("4.3. De la note à la décision")
para(
    "Le score de risque est converti en probabilité par une fonction sigmoïde "
    "centrée sur 50, puis transformé en décision selon trois bandes : en dessous "
    "de 45, la transaction est APPROUVÉE ; entre 45 et 70, elle est marquée "
    "SUSPECTE et placée en examen ; au-delà de 70, elle est BLOQUÉE (et qualifiée "
    "de risque critique au-delà de 85). Le moteur produit simultanément des "
    "valeurs de type SHAP et les cinq principales raisons en langage naturel, qui "
    "constituent la couche d'explication présentée à l'utilisateur."
)

h2("4.4. L'architecture de streaming et le déroulement de la simulation")
para(
    "La simulation s'exécute dans une boucle asynchrone (asyncio) pilotée par un "
    "service dédié. À chaque itération, un lot de transactions est généré au "
    "rythme paramétré (10, 100 ou 1 000 transactions par seconde), chaque "
    "transaction est notée, le résultat est persisté en base via la couche CRUD, "
    "puis publié sur un broker d'événements en mémoire. Le broker rediffuse "
    "l'événement à tous les abonnés (le tableau de bord) par Server-Sent Events "
    "ou WebSocket, avec un tampon de rejeu pour les nouveaux clients. Chaque "
    "session de simulation ouvre par ailleurs un enregistrement « SimulationRun » "
    "permettant à Power BI de découper l'analyse par exécution. La Figure 4 "
    "détaille cette séquence."
)
add_figure(os.path.join(FIG, "fig4_sequence_streaming.png"),
           "Séquence d'exécution d'un cycle de la simulation en temps réel.")

page_break()

# =====================================================================
# 5. ARCHITECTURE TECHNIQUE RÉELLE
# =====================================================================
h1("5. Architecture technique réellement développée")
para(
    "Cette section décrit l'architecture telle qu'elle a été effectivement "
    "implémentée et conteneurisée : un backend FastAPI, une base de données "
    "PostgreSQL pilotée par un ORM SQLAlchemy, un frontend web servi par l'API, le "
    "tout orchestré par Docker Compose, et une couche d'intégration Power BI en "
    "DirectQuery. La Figure 5 en donne la vue de déploiement."
)
add_figure(os.path.join(FIG, "fig7_tech_architecture.png"),
           "Architecture technique réelle : conteneurs Docker, backend, base de "
           "données, frontend et intégration Power BI.")

h2("5.1. Conteneurisation et déploiement (Docker)")
para(
    "Le système est décrit par un fichier docker-compose qui orchestre deux "
    "conteneurs. Le service api est construit à partir de l'image "
    "python:3.11-slim (Dockerfile.api) : il installe les dépendances de "
    "production, copie le code source, les configurations, les scripts et le "
    "magasin de modèles, puis démarre le serveur. Le service db est une instance "
    "PostgreSQL 16 officielle, dont les données sont persistées dans un volume "
    "nommé. Le conteneur api reçoit l'URL de la base par variable "
    "d'environnement et monte le répertoire des modèles, ce qui découple le code "
    "des artefacts d'apprentissage."
)
add_table(
    ["Service", "Image / build", "Port", "Rôle"],
    [
        ["api", "Dockerfile.api (python:3.11-slim)", "8080", "Backend FastAPI + frontend servi à /dashboard."],
        ["db", "postgres:16", "5432", "Base PostgreSQL (volume persistant fraud_pgdata)."],
    ],
    widths=[0.9, 2.6, 0.7, 2.3], font=9,
)
para(
    "L'image d'API repose sur des dépendances volontairement réduites "
    "(requirements-api.txt) : FastAPI, Uvicorn et Pydantic pour le service web ; "
    "SQLAlchemy et psycopg2 pour la base ; scikit-learn, XGBoost, LightGBM, "
    "CatBoost et imbalanced-learn pour l'inférence des modèles ; Loguru pour la "
    "journalisation. Le calcul SHAP est désactivé par défaut dans l'image "
    "(variable ENABLE_SHAP=0) pour alléger le démarrage, et réactivable à la "
    "demande.",
    italic=True,
)

h2("5.2. Le backend FastAPI et les deux moteurs de décision")
para(
    "Le backend est une application FastAPI servie par Uvicorn. Au démarrage, un "
    "gestionnaire de cycle de vie initialise la base (création des tables et "
    "migration additive des colonnes manquantes). L'application active une "
    "politique CORS, sert le frontend à la route /dashboard, et enregistre "
    "plusieurs routeurs de façon défensive — une panne d'un routeur secondaire ne "
    "peut pas faire tomber le cœur de prédiction."
)
para(
    "Un point essentiel de l'architecture réelle est la coexistence de deux "
    "moteurs de décision complémentaires :"
)
bullet("il charge un modèle supervisé entraîné (par défaut l'ensemble configuré) "
       "avec son préprocesseur et son ingénieur de variables, applique exactement "
       "le même pipeline qu'à l'entraînement, produit une probabilité puis une "
       "décision, et fournit une explication SHAP (TreeExplainer avec repli). Il "
       "sert le point d'entrée /api/v1/predict ;", "Le moteur ML hors-ligne (FraudDecisionEngine) : ")
bullet("c'est le moteur additif explicable décrit en section 4. Il évalue les "
       "transactions générées par le simulateur et alimente, sous l'étiquette "
       "model_version « dynamic-fraud-engine », les pages de supervision temps "
       "réel du tableau de bord et de Power BI.", "Le moteur dynamique (DynamicFraudScoringEngine) : ")
para(
    "Autrement dit, les pages « live » du tableau de bord reflètent le moteur "
    "dynamique, tandis que les pages de benchmark, de coût et de seuil reflètent "
    "les modèles supervisés évalués hors-ligne. Le tableau ci-dessous récapitule "
    "les principaux points d'entrée réellement exposés."
)
add_table(
    ["Point d'entrée", "Fonction"],
    [
        ["POST /api/v1/predict", "Noter une transaction avec le modèle ML hors-ligne."],
        ["POST /predict", "Noter une transaction bancaire avec le moteur dynamique."],
        ["GET /transactions, /customers, /fraud-alerts", "Lister transactions, clients et alertes."],
        ["GET /analytics/live", "Analytique de fraude agrégée en direct."],
        ["GET /shap-explanation", "Explication SHAP-style d'une transaction."],
        ["POST /simulator/start | /stop ; GET /status", "Piloter et suivre le simulateur."],
        ["GET /live-stream (SSE) ; /ws/live-stream (WebSocket)", "Diffuser le flux d'événements en temps réel."],
        ["GET /api/powerbi/* ; /api/v1/export/*", "Endpoints Power BI (JSON) et exports CSV/ZIP."],
        ["GET /health, /docs, /dashboard", "Sonde, documentation Swagger, frontend."],
    ],
    widths=[3.1, 3.4], font=8.7,
)

h2("5.3. La base de données PostgreSQL et le schéma SQLAlchemy")
para(
    "La persistance repose sur SQLAlchemy. L'URL de connexion est lue dans une "
    "variable d'environnement (avec réécriture automatique du préfixe "
    "postgres:// hérité de certains hébergeurs) : PostgreSQL est la cible de "
    "production et de Power BI DirectQuery, tandis que SQLite sert de repli local "
    "pour les tests. L'initialisation crée les tables et applique une migration "
    "additive légère qui ajoute les colonnes manquantes aux bases de "
    "démonstration existantes. Les schémas réellement définis se répartissent en "
    "trois familles."
)
add_table(
    ["Famille de tables", "Tables (réelles)", "Usage"],
    [
        ["Cœur /predict (historique)", "transactions, predictions, fraud_alerts", "Chemin de prédiction ML classique."],
        ["Plateforme bancaire live", "banking_customers, banking_transactions, banking_alerts, banking_shap_explanations", "Transactions, alertes et explications du moteur dynamique."],
        ["Couche analytique Power BI", "simulation_runs, model_benchmark_results, cost_comparison_summary, cost_analysis_sweep, powerbi_exports", "Faits de benchmark/coût/seuil chargés depuis les CSV + suivi des runs."],
    ],
    widths=[1.8, 2.7, 2.0], font=8.3,
)
para(
    "Les tables analytiques ne sont jamais écrites par les pipelines "
    "d'entraînement ou de simulation : elles sont alimentées de façon idempotente "
    "depuis les artefacts CSV précalculés (script de chargement de l'entrepôt "
    "Power BI), les CSV restant la source de vérité. Pour la production, "
    "l'introduction d'un outil de migration de schéma (Alembic) constitue une "
    "perspective d'amélioration ; à ce stade, la migration additive intégrée "
    "suffit aux besoins de démonstration.",
    italic=True,
)

h2("5.4. Le frontend web")
para(
    "Le frontend est une application web servie directement par l'API à la route "
    "/dashboard (« Atlas Fraud Intelligence Platform »). Il est réalisé en un "
    "seul fichier autonome s'appuyant sur React 18 (distribution UMD) et Babel "
    "Standalone — qui compile le JSX dans le navigateur —, complété par Chart.js "
    "pour les graphiques, Leaflet pour la carte interactive et Lucide pour les "
    "icônes, sur un thème sombre « FinTech ». Il consomme les points d'entrée de "
    "la plateforme bancaire et s'abonne au flux temps réel (SSE / WebSocket) pour "
    "actualiser en continu les tables, la carte de fraude et les alertes pendant "
    "qu'une simulation tourne. Ce frontend intégré est complémentaire du rapport "
    "Power BI : le premier vise la supervision opérationnelle immédiate, le "
    "second l'analyse décisionnelle approfondie."
)

h2("5.5. L'intégration Power BI")
para(
    "L'intégration analytique est traitée en détail à la section 7. Du point de "
    "vue technique, elle repose sur neuf vues SQL en lecture seule (vw_powerbi_*) "
    "interrogées par Power BI en DirectQuery sur PostgreSQL ; un jeu d'endpoints "
    "/api/powerbi/* fournit une voie de repli en import JSON, et un endpoint "
    "d'export produit un paquet (ZIP) de fichiers prêts pour un usage hors "
    "connexion. Les vues additives n'altèrent jamais les tables opérationnelles."
)

# =====================================================================
# 6. PROTOCOLE D'ÉVALUATION (résultats définitifs en attente)
# =====================================================================
h1("6. Protocole d'évaluation et indicateurs")
p = para(
    "Note importante : conformément à la démarche adoptée, les résultats chiffrés "
    "définitifs (tableaux, métriques, graphiques et exports) seront intégrés ici "
    "à réception des livrables finaux. La présente section définit le protocole "
    "d'évaluation réellement utilisé et les indicateurs retenus pour produire et "
    "analyser ces résultats, sans préjuger des valeurs finales.",
    bold=True,
)
para(
    "L'analyse définitive portera sur huit volets : les performances des modèles, "
    "le choix du meilleur modèle, les faux positifs et faux négatifs, le seuil "
    "optimal, les coûts, l'explicabilité SHAP, la simulation et la restitution "
    "Power BI. Le protocole de chacun est précisé ci-dessous."
)

h2("6.1. Indicateurs de performance des modèles")
para(
    "Les modèles sont évalués sur le hold-out chronologique. En raison du "
    "déséquilibre extrême (0,17 % de fraudes), la métrique de référence est la "
    "PR-AUC, complétée par le coefficient de corrélation de Matthews (MCC) : "
    "toutes deux résistent au déséquilibre, contrairement à l'exactitude et, dans "
    "une moindre mesure, à la ROC-AUC. Les indicateurs calculés pour chaque "
    "modèle sont les suivants."
)
add_table(
    ["Indicateur", "Définition / rôle dans l'analyse"],
    [
        ["Précision", "Part de vraies fraudes parmi les transactions bloquées."],
        ["Rappel (sensibilité)", "Part des fraudes effectivement détectées."],
        ["F1", "Moyenne harmonique précision/rappel."],
        ["MCC", "Corrélation prédiction/vérité, robuste au déséquilibre."],
        ["PR-AUC", "Métrique de référence (aire précision-rappel)."],
        ["ROC-AUC", "Pouvoir discriminant global (complémentaire)."],
        ["Spécificité, exactitude équilibrée", "Contrôle du comportement sur la classe majoritaire."],
        ["Matrice de confusion (TP/FP/TN/FN)", "Base du décompte des erreurs et du coût."],
    ],
    widths=[2.5, 4.0], font=9,
)

h2("6.2. Choix du meilleur modèle")
para(
    "Le meilleur modèle n'est pas choisi a priori : le classement est piloté par "
    "les données et combine la performance technique et, lorsque les coûts sont "
    "chargés, la dimension économique. Le modèle retenu est marqué par un "
    "indicateur dédié dans la couche analytique. L'analyse finale justifiera ce "
    "choix sous cinq angles : technique, économique, opérationnel, aptitude au "
    "déploiement et limites."
)

h2("6.3. Faux positifs et faux négatifs")
para(
    "Les deux types d'erreur n'ont pas le même coût : un faux négatif (fraude "
    "manquée) entraîne une perte financière proportionnelle au montant, tandis "
    "qu'un faux positif (blocage légitime) engendre des coûts d'investigation, de "
    "marge perdue et d'insatisfaction. L'analyse comparera, par modèle et par "
    "seuil, le nombre de FP et de FN et leur traduction monétaire."
)

h2("6.4. Seuil optimal de décision")
para(
    "Le seuil n'est pas figé à 0,50. Un balayage complet est effectué de 0,01 à "
    "0,99 par pas de 0,01 ; pour chaque seuil sont calculés les métriques et le "
    "coût total. Trois seuils optimaux sont identifiés — selon F1, selon MCC et "
    "selon le coût — et repérés par des indicateurs dédiés. L'analyse finale "
    "comparera le seuil par défaut au seuil optimal et chiffrera l'économie "
    "associée."
)

h2("6.5. Coûts")
para(
    "L'évaluation économique repose sur la fonction de coût asymétrique "
    "C(s) = FN(s)·C_FN + FP(s)·C_FP, en mode « sensible au montant » dont la "
    "composition a été détaillée en section 3.8. L'analyse finale présentera le "
    "coût total par modèle et par stratégie de seuil, ainsi que l'économie "
    "réalisée par rapport au seuil de référence."
)

h2("6.6. Explicabilité (SHAP)")
para(
    "Deux dispositifs d'explication coexistent. Pour les modèles supervisés, un "
    "explainer SHAP (TreeExplainer, avec repli) attribue à chaque variable sa "
    "contribution. Pour le moteur dynamique, les contributions additives sont "
    "calculées et persistées par transaction, puis traduites en libellés métier. "
    "L'analyse finale présentera les variables globalement les plus contributives "
    "et des exemples de décompositions locales."
)

h2("6.7. Simulation")
para(
    "Chaque exécution du simulateur est tracée (enregistrement de run) avec ses "
    "paramètres (rythme, ratio de fraude, version du moteur) et ses métriques "
    "agrégées : transactions générées, fraudes simulées, détectées et manquées, "
    "faux positifs, taux de détection, précision sur les transactions bloquées, "
    "rappel sur la fraude et taux de faux positifs. L'analyse finale exploitera "
    "ces métriques pour valider l'efficacité du dispositif, y compris par "
    "scénario de fraude."
)

h2("6.8. Restitution Power BI")
para(
    "Les indicateurs ci-dessus sont restitués dans le rapport Power BI décrit en "
    "section 7. L'analyse finale s'appuiera sur les KPI réellement présentés "
    "(volumes, taux de fraude, pertes évitées, taux de pertes évitées, coût des "
    "faux positifs, classement des modèles, seuils optimaux, contributions SHAP).",
)

page_break()

# =====================================================================
# 7. POWER BI DASHBOARD — PAGE BY PAGE
# =====================================================================
h1("7. Le tableau de bord Power BI : structure réelle, page par page")
para(
    "Le rapport Power BI réellement développé compte neuf pages (telles "
    "qu'observées dans le tableau de bord livré). Il est connecté en DirectQuery "
    "à la base PostgreSQL, qu'il interroge à travers des vues analytiques en "
    "lecture seule (vw_powerbi_*), sans copier ni altérer les données "
    "opérationnelles. Cette section décrit la structure réelle de chaque page — "
    "ses indicateurs, ses visuels et sa source — l'analyse chiffrée définitive "
    "étant traitée à la section 6 à réception des exports finaux. Les valeurs "
    "citées ci-après proviennent de l'instance de démonstration et illustrent la "
    "structure, sans constituer le résultat définitif."
)

h2("7.1. Modèle de données et conventions de restitution")
para(
    "Le rapport est organisé en modèle en étoile (Figure 6) : la table de faits "
    "centrale (transactions live) est reliée aux dimensions conformes (date, "
    "client, statut de décision, niveau de risque) ainsi qu'aux faits d'alertes "
    "et d'explications SHAP par l'identifiant de transaction. Les tables de grain "
    "différent (benchmark des modèles, analyse de coût, balayage de seuil, "
    "contexte marocain, KPI instantanés) restent volontairement non reliées et "
    "sont filtrées par leurs propres segments : cette séparation évite tout "
    "double comptage, notamment entre transactions et explications SHAP."
)
add_figure(os.path.join(FIG, "fig5_powerbi_star_schema.png"),
           "Modèle en étoile du rapport Power BI et grain des tables de faits.")
para(
    "Quatre conventions garantissent l'honnêteté méthodologique de la "
    "restitution, essentielles à la défense du travail :"
)
bullet("la probabilité de fraude est exprimée sur une échelle 0–100 (moteur "
       "dynamique), le seuil de blocage exposé étant fixé à 70 ;", "Échelle de score : ")
bullet("le statut de décision réellement émis prend trois valeurs — APPROVED, "
       "SUSPICIOUS et BLOCKED ; la bande REVIEW est prévue dans le modèle mais "
       "n'est pas encore produite par le moteur dynamique ;", "Décisions : ")
bullet("le label de fraude est la vérité-terrain de la simulation (scénario ≠ "
       "« normal ») ; dans une banque réelle, il ne serait pas disponible au "
       "moment de la décision — le rapport le signale explicitement ;", "Vérité-terrain : ")
bullet("les données marocaines sont synthétiques et porteuses d'un avertissement "
       "obligatoire (« non représentatives de la fraude bancaire réelle au "
       "Maroc »).", "Transparence : ")

h2("7.2. Cartographie des neuf pages")
para(
    "Les neuf pages se répartissent en cinq thématiques (Figure 7) : la "
    "supervision temps réel, l'analyse économique, la comparaison des modèles, "
    "l'explicabilité et le contexte de déploiement. Le Tableau ci-dessous en "
    "donne la synthèse."
)
add_figure(os.path.join(FIG, "fig6_powerbi_pages_map.png"),
           "Cartographie thématique des neuf pages réelles du rapport Power BI.")
add_table(
    ["Page", "Thématique", "Contenu principal"],
    [
        ["P1 — Monitoring exécutif", "Temps réel", "KPI globaux, anneau des décisions, carte, courbe temporelle."],
        ["P2 — Transactions live", "Temps réel", "Journal détaillé, nuage montant/risque, barres par niveau de risque."],
        ["P3 — Alertes", "Temps réel", "Alertes par sévérité, par heure, par ville ; compteurs d'alertes."],
        ["P4 — Pertes évitées & coût", "Économique", "Waterfall des pertes, montants par décision, taux de pertes évitées."],
        ["P5 — Benchmark des modèles", "Modèles", "Tableau des 7 modèles, F1/precision/recall, performance vs latence."],
        ["P6 — Coût & économies par modèle", "Économique", "cost_saved, coût FP/FN, coût total et économies par modèle."],
        ["P7 — Optimisation du seuil", "Économique", "MCC, F1 et coût total en fonction du seuil ; seuils recommandés."],
        ["P8 — Explicabilité SHAP", "Explicabilité", "Contributions par variable pour la transaction sélectionnée."],
        ["P9 — Contexte marocain", "Contexte", "Volume et taux de fraude simulé par ville (+ disclaimer)."],
    ],
    widths=[2.2, 1.3, 3.0], font=8.4,
)

h2("7.3. Pages de supervision temps réel (P1, P2, P3)")
h3("Page 1 — Monitoring exécutif")
para(
    "Cette page offre la vue instantanée de la plateforme. Une bande de cartes "
    "KPI présente le nombre total de transactions, les volumes approuvés, suspects "
    "et bloqués, le taux de fraude et les pertes évitées (dans l'instance de "
    "démonstration : 48 K transactions, ~60 % approuvées, ~19 % bloquées, ~21 % "
    "suspectes, taux de fraude 12,86 %). Les visuels associent un anneau de "
    "répartition des décisions (decision_status), une jauge de taux de fraude, "
    "une carte mondiale des transactions détectées (latitude/longitude) et une "
    "courbe de volume par horodatage."
)
h3("Page 2 — Transactions live")
para(
    "Le visuel central est un tableau exhaustif de chaque transaction générée "
    "(identifiant, horodatage, client, montant, score de risque, décision, niveau "
    "de risque, ville, version du moteur). Une mise en forme conditionnelle colore "
    "la décision (vert / orange / rouge) et le score de risque selon les bandes "
    "45 et 70. Un nuage de points croise le montant et le score de risque, et un "
    "histogramme répartit les transactions par niveau de risque (Low / Watch / "
    "High). Des cartes rappellent le total, le montant total traité (~24,81 M) et "
    "le taux de fraude."
)
h3("Page 3 — Alertes opérationnelles")
para(
    "Cette page mesure la pression opérationnelle. Un tableau liste les alertes "
    "avec leur sévérité (CRITICAL, MEDIUM), leur score de risque, leur montant, "
    "leur ville et leur décision. Des histogrammes présentent les alertes "
    "générées par heure et par niveau de risque, ainsi que le nombre d'alertes par "
    "ville. Deux compteurs synthétisent le nombre d'alertes et les alertes "
    "ouvertes (≈ 19 K dans l'instance de démonstration). À ce stade, toute alerte "
    "est considérée « ouverte » : un véritable flux d'acquittement analyste "
    "constitue une perspective (cf. 7.8)."
)

h2("7.4. Pages d'analyse économique (P4, P7)")
h3("Page 4 — Pertes évitées et analyse de coût")
para(
    "Cette page matérialise le cadre de coût asymétrique. Les cartes présentent le "
    "montant total traité, le montant total des fraudes, les pertes évitées, le "
    "taux de pertes évitées (96,65 % dans l'instance de démonstration) et le coût "
    "des faux positifs. Un graphique en cascade (waterfall) décompose l'impact "
    "financier — fraudes détectées, perte résiduelle, pertes évitées, total — et "
    "un histogramme compare les montants approuvés, suspects et bloqués par "
    "statut de décision."
)
h3("Page 7 — Optimisation du seuil")
para(
    "Alimentée par le balayage complet des seuils, cette page montre comment le "
    "seuil gouverne la performance technique et le coût. Des courbes tracent le "
    "MCC, le F1 et le coût total en fonction de la valeur de seuil, par modèle. Un "
    "tableau « Seuils recommandés par critère d'optimisation » récapitule, pour "
    "chaque modèle, les seuils optimaux selon F1, MCC et coût (indicateurs "
    "is_f1_optimal, is_mcc_optimal, is_cost_optimal)."
)

h2("7.5. Pages de comparaison des modèles (P5, P6)")
h3("Page 5 — Benchmark des modèles")
para(
    "Cette page compare les sept modèles entraînés à partir des valeurs réellement "
    "chargées (aucune métrique inventée). Des cartes mettent en avant les "
    "meilleurs scores (meilleur rappel, meilleur F1, meilleure précision) ; un "
    "tableau présente, par modèle, exactitude, précision, rappel, F1, MCC et "
    "ROC-AUC ; un nuage de points croise le F1 et le temps d'inférence "
    "(performance vs latence) ; des barres comparent le compromis "
    "précision/rappel et le F1 par modèle."
)
h3("Page 6 — Coût et économies par modèle")
para(
    "Cette page traduit le benchmark en termes économiques. Un tableau présente, "
    "par modèle, l'économie réalisée (cost_saved), le coût des faux positifs, le "
    "coût des faux négatifs et le coût total. Des cartes synthétisent le coût "
    "total estimé et les économies générées ; des histogrammes classent les "
    "modèles par économie réalisée et confrontent économies et coût des faux "
    "positifs."
)

h2("7.6. Page d'explicabilité (P8)")
para(
    "La page SHAP explique pourquoi une transaction est bloquée. Un segment "
    "permet de sélectionner une transaction ; des cartes affichent son montant, "
    "son score de risque, sa probabilité de fraude et sa décision. Un tableau "
    "« Explication technique et formulation métier » présente, par variable, la "
    "somme des valeurs SHAP et la formulation associée (valeur observée et sens de "
    "la contribution), tandis qu'un histogramme classe les variables ayant "
    "influencé la décision (par exemple amount_vs_customer_average, "
    "account_age_days, merchant_category, geographic_risk, ip_reputation, "
    "impossible_travel, transaction_frequency_60s). La page distingue clairement "
    "les contributions SHAP brutes de leur traduction en libellé métier."
)

h2("7.7. Page de contexte et d'applicabilité (P9)")
para(
    "La dernière page discute l'adaptation au marché marocain à partir des "
    "transactions synthétiques localisées au Maroc. Un histogramme présente le "
    "volume de transactions par ville (Marrakech, Casablanca, Agadir, Fès, Rabat, "
    "Tanger), un autre le taux de fraude simulé par ville, et un tableau "
    "récapitule par ville le nombre de transactions, le montant total et le taux "
    "de fraude. Une note méthodologique visible rappelle obligatoirement que ces "
    "données sont synthétiques et non représentatives de la fraude bancaire réelle "
    "au Maroc."
)

h2("7.8. Fonctionnalités en cours et perspectives")
para(
    "Pour rester fidèle à l'état réel du projet, certaines fonctionnalités "
    "documentées dans les spécifications ne sont pas (encore) présentes dans le "
    "rapport livré et sont présentées comme perspectives :"
)
bullet("une page dédiée à l'efficacité de la simulation par exécution (matrice "
       "de confusion et taux de détection par run), dont les données existent "
       "déjà dans la table des runs ;", "Page « efficacité de simulation » : ")
bullet("une page de détail par transaction (drill-through) reliant un cas "
       "individuel à sa décomposition SHAP et à sa fiche client ;", "Drill-through : ")
bullet("la bande de décision REVIEW (examen humain), modélisée mais non encore "
       "émise par le moteur dynamique ;", "Bande REVIEW : ")
bullet("un flux d'acquittement des alertes (ouvert / pris en charge / clôturé), "
       "le modèle ne disposant pour l'instant que d'un état implicite "
       "« ouvert ».", "Workflow d'alertes : ")

page_break()

# =====================================================================
# 8. DISCUSSION
# =====================================================================
h1("8. Discussion critique et perspectives")
para(
    "Plusieurs limites et précautions doivent accompagner ces résultats."
)
h3("8.1. Biais du jeu de données et généralisation")
para(
    "Le jeu européen est anonymisé par ACP, ce qui prive d'interprétation "
    "métier, et limité à 48 heures, ce qui interdit de modéliser la dérive "
    "conceptuelle et la saisonnalité. Surtout, les signatures de fraude "
    "européennes (puce-et-code, escalade en card-not-present) diffèrent de "
    "celles d'autres écosystèmes — notamment marocain (mixte bande "
    "magnétique/EMV, fraudes mobiles, interception d'OTP, SIM-swap). Un modèle "
    "entraîné uniquement sur ces données risque donc de sur-apprendre des "
    "signatures locales et de manquer des fraudes typiques d'ailleurs, ce qui "
    "plaide pour l'apprentissage par transfert et l'apprentissage fédéré dans un "
    "déploiement réel."
)
h3("8.2. Explicabilité et conformité")
para(
    "Une ROC-AUC de 0,99 n'a aucune valeur sans justification par décision : "
    "bloquer un paiement constitue une décision automatisée à effet significatif, "
    "soumise à des obligations d'explication (RGPD, futur règlement européen sur "
    "l'IA ; au Maroc, loi 09-08 et lignes directrices de la CNDP). Le système y "
    "répond en couplant des attributions de type SHAP à des gabarits de "
    "narration, afin que l'utilisateur reçoive des raisons en langage clair "
    "plutôt que des coefficients bruts."
)
h3("8.3. Réduction des faux positifs et enjeux éthiques")
para(
    "Chaque blocage erroné a un coût (revenu perdu, appels au support, risque "
    "d'attrition) et peut laisser un client légitime sans moyen de paiement à "
    "l'étranger. Quatre mécanismes le limitent : calibration des probabilités, "
    "optimisation du seuil par le coût, décision en trois bandes avec mise en "
    "examen humaine, et filtrage par ensemble. Sur le plan éthique, "
    "l'anonymisation des variables masque tout recours éventuel à un proxy "
    "d'attribut protégé : un audit d'équité sur les variables brutes est "
    "indispensable avant toute mise en production."
)
h3("8.4. Dérive conceptuelle")
para(
    "Les schémas de fraude évoluent (card-testing, attaques BIN, interception "
    "d'OTP, prise de contrôle par SIM-swap). Le modèle doit être réentraîné sur "
    "fenêtre glissante et surveillé (indice de stabilité de population, suivi des "
    "variables SHAP dominantes, tests champion/challenger). Les points "
    "d'extension nécessaires sont déjà exposés par les points d'entrée de "
    "monitoring de l'API."
)
h3("8.5. Perspectives et éléments en cours de développement")
para(
    "Plusieurs évolutions sont identifiées comme perspectives, distinctes de "
    "l'existant réellement implémenté :"
)
bullet("introduction d'un outil de migration de schéma (Alembic) en remplacement "
       "de la migration additive intégrée, pour un usage en production ;", "Migrations de base : ")
bullet("émission effective de la bande de décision REVIEW par le moteur "
       "dynamique et mise en place d'un flux d'acquittement des alertes ;", "Examen humain : ")
bullet("ajout au rapport Power BI d'une page d'efficacité de simulation par "
       "exécution et d'une page de détail par transaction (drill-through) ;", "Restitution : ")
bullet("activation du suivi des expériences (MLflow) et de tests "
       "champion/challenger pour le réentraînement périodique ;", "MLOps : ")
bullet("constitution d'un jeu de données marocain anonymisé et recours à "
       "l'apprentissage par transfert ou fédéré pour la généralisation locale.", "Données locales : ")

# =====================================================================
# 9. SYNTHÈSE
# =====================================================================
h1("9. Synthèse")
para(
    "La partie empirique a établi une chaîne de bout en bout, reproductible et "
    "auditable, articulée autour de deux volets complémentaires. Le volet "
    "hors-ligne fournit un pipeline d'apprentissage à garanties anti-fuite "
    "strictes, un banc d'essai de sept modèles et de plusieurs stratégies de "
    "rééquilibrage, une calibration et une optimisation de seuil orientées coût. "
    "Le volet temps réel transpose ces principes dans une simulation bancaire "
    "réaliste, fondée sur une génération synthétique à deux niveaux (clients puis "
    "transactions) et un moteur de scoring additif explicable, diffusant des "
    "décisions traçables en continu. Ces deux volets s'appuient sur une "
    "architecture technique réellement développée et conteneurisée — backend "
    "FastAPI, base PostgreSQL pilotée par SQLAlchemy, frontend web temps réel et "
    "intégration Power BI en DirectQuery — décrite en section 5."
)
para(
    "Le choix du meilleur modèle, les performances chiffrées, l'analyse des faux "
    "positifs et négatifs, le seuil optimal, les coûts, l'explicabilité SHAP et "
    "l'efficacité de la simulation seront établis de manière définitive en "
    "section 6, sur la base des tableaux, métriques, graphiques et exports "
    "finaux. La méthodologie, l'architecture et le protocole étant désormais "
    "fixés, l'intégration de ces résultats viendra clore la démonstration de "
    "faisabilité d'un système antifraude performant, économiquement justifié et "
    "conforme aux exigences d'explicabilité."
)

try:
    doc.save(OUT)
    print("SAVED:", OUT)
except PermissionError:
    alt = OUT.replace(".docx", "_v2.docx")
    doc.save(alt)
    print("LOCKED original (open in Word?). SAVED instead:", alt)
print("Figures inserted:", _fig_counter["n"])
